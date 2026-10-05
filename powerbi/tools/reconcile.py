#!/usr/bin/env python3
"""
Reconcile the Power BI model's logic against the web dashboard.

Usage:
    python powerbi/tools/reconcile.py <case_export.xlsx> <survey_analysis.xlsx> [web_numbers.json]

A line-by-line Python port of the model's Power Query (powerbi/tools/m/*.m) and
DAX measures (measures.py), run on the same files the web dashboard was built
from. Prints the figures the Power BI pages should show with no slicers set;
when the web dashboard's own figures are given (from its compute code), they
are compared side by side. Use the printed table to check the Power BI report
after the first refresh.
"""
import datetime
import json
import math
import os
import re
import sys
import warnings

warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build_model import literal_queries  # noqa: E402

EPOCH = datetime.datetime(1899, 12, 30)


def oad(d):
    return None if d is None else (d - EPOCH).total_seconds() / 86400


def rnd(x):
    return None if x is None else math.floor(x + 0.5)


def mtext(v):
    return '' if v is None else str(v)


def collapse(v):
    return '' if v is None else ' '.join(re.split(r'[ \t\n\r\xa0]+', str(v))).strip()


# ---- lookups, parsed back out of the generated M literals so the port uses the same data
L = literal_queries()


def rows_of(mtable):
    body = mtable[mtable.index('],'):]       # skip the column types; each row is an innermost { … }
    return [re.findall(r'"((?:[^"]|"")*)"|(-?\d+(?:\.\d+)?)', r) for r in re.findall(r'\{([^{}]*)\}', body)]


def text_rows(name):
    out = []
    for r in rows_of(L[name]):
        out.append([a.replace('""', '"') if a or not b else float(b) for a, b in r])
    return out


REGION_REF = text_rows('RegionReference')
OFFER = {k: v for k, v in text_rows('OfferLookup')}
HQ = {k: v for k, v in text_rows('HqOffices')}
EXCLUDED_RES = re.findall(r'"([^"]+)"', L['ExcludedResolutions'].split('\n', 1)[1])
MANUAL = {r[0]: r[1:] for r in text_rows('ManualCoding')}
SATFIX = {r[0]: (r[1], r[2]) for r in text_rows('SatisfactionCorrections')}
EXCL_RESP = re.findall(r'"([^"]+)"', L['ExcludedResponses'].split('\n', 1)[1])
COORDS = {r[0]: (r[1], r[2]) for r in text_rows('OfficeCoords')}


def norm_office(raw):                      # fnNormOffice
    s = (raw or '').lower()
    out, depth = '', 0
    for c in s:
        if c == '(':
            depth += 1
        elif c == ')' and depth > 0:
            out += ' '; depth -= 1
        elif depth > 0:
            pass
        else:
            out += c
    out = ''.join(' ' if c in ".,'-" else c for c in out)
    toks = [t for t in re.split(r'[ \t\n\r\xa0]', out) if t]
    filler = {'the', 'of', 'and', 'rep', 'republic', 'dem', 'state', 'united', 'peoples'}
    kept = []
    for i, t in enumerate(toks):
        prev = toks[i - 1] if i else ''
        nxt = toks[i + 1] if i < len(toks) - 1 else ''
        if t in filler or (t == 'people' and nxt == 's') or t == 'island' or (t == 's' and prev in ('people', 'island')):
            continue
        kept.append(t)
    return ' '.join(kept)


LOOKUP = {}
for name, office, region in REGION_REF:    # OfficeLookup: later rows win
    LOOKUP[norm_office(name)] = (office, region)


def map_office(raw):                       # fnMapOffices
    raw = (raw or '').strip()
    if raw in HQ:
        return HQ[raw], 'HQ'
    return LOOKUP.get(norm_office(raw), (raw, 'Unmapped'))


def map_offer(raw):                        # fnMapOffer
    s = mtext(raw)
    key = ' '.join(t for t in re.split(r'[ \t\n\r]', s.replace('\xa0', ' ').replace('-', ' ').lower()) if t)
    return OFFER.get(key, s.strip())


def map_practice(raw):                     # fnMapPractice
    s = mtext(raw)
    return 'Programme Policy & Strategy' if collapse(s.lower()) in ('other', 'innovation') else s.strip()


def map_type(raw):                         # fnMapType
    s = mtext(raw)
    k = collapse(s.lower())
    if k in ('big ticket', 'big ticket item', 'big-ticket'):
        return 'Big Ticket'
    if k in ('routine', 'regular'):
        return 'Routine'
    return s.strip()


ENT = [('&nbsp;', ' '), ('&#160;', ' '), ('&lt;', '<'), ('&gt;', '>'), ('&quot;', '"'), ('&#39;', "'"), ('&apos;', "'"),
       ('&rsquo;', '’'), ('&lsquo;', '‘'), ('&rdquo;', '”'), ('&ldquo;', '“'), ('&ndash;', '–'), ('&mdash;', '—'), ('&amp;', '&')]


def clean_html(v):                          # fnCleanHtml
    if v is None:
        return ''
    parts = str(v).split('<')
    s = parts[0] + ''.join(' ' + (p.split('>', 1)[1] if '>' in p else '') for p in parts[1:])
    for a, b in ENT:
        s = s.replace(a, b)
    return collapse(s)


EXACT = {'', 'na', 'n/a', 'n.a.', 'n', 'nil', 'none', 'undefined', 'missing', 'tbd', 'tba', 'test', 'testing', 'tests',
         '.', '-', '--', '...', 'x', 'xx', 'xxx', '?', 'pending'}
PREFIXES = ('please add', 'add description', 'add decription', 'add descripton', 'please provide more', 'to be added',
            'to be defined', 'to be confirmed', 'description to follow', 'test ', 'testing ')


def placeholder(t):                         # fnIsPlaceholder
    n = t.lower().strip().rstrip('.!:;,').strip()
    return n in EXACT or n.startswith(PREFIXES)


def load_requests(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        rec = dict(zip(head, r))
        for k in ('Expected Start Date', 'Expected Completion Date', 'Created', 'Opened', 'Updated', 'Resolved', 'Closed'):
            v = rec.get(k)
            rec[k] = v if isinstance(v, datetime.datetime) else None
        for k in head:
            if not isinstance(rec[k], datetime.datetime) and rec[k] is not None:
                rec[k] = str(rec[k])
        for k in head:
            if rec[k] is None and k not in ('Expected Start Date', 'Expected Completion Date', 'Created', 'Opened', 'Updated', 'Resolved', 'Closed'):
                rec[k] = ''
        rec['Request ID'] = rec['Case Report'] or rec['Number']
        if rec['Request ID']:
            out.append(rec)
    return out


def requests_table(raw):                    # Requests.m
    kept = [r for r in raw if r['Resolution code'] not in EXCLUDED_RES]
    stamps = [oad(r[k]) for r in kept for k in ('Created', 'Opened', 'Updated') if r[k] is not None]
    today = math.floor(max(stamps))
    real = {'ESAR', 'APR', 'WCAR', 'LACR', 'ECAR', 'MENAR'}
    rows = []
    for i, r in enumerate(kept):
        office, region = map_office(r['Office/Division'])
        st, lead = r['Implementation Status'], r['Assigned to']
        xs, xc = oad(r['Expected Start Date']), oad(r['Expected Completion Date'])
        received = r['Created'] or r['Opened']
        details = clean_html(r['Details/Description'])
        ph = placeholder(details)
        hd = not ph
        is_co = region in real
        closedish = st in ('100%', 'Discontinued')
        past = xc is not None and xc < today
        overdue = not closedish and past
        late = None if xc is None else today - xc
        active = st not in ('100%', 'Discontinued', 'Unassigned')
        in_review = st in ('Unassigned', '0%')
        wait_from = received if st == 'Unassigned' else (r['Updated'] or received)
        waiting = None if wait_from is None else rnd(today - oad(wait_from))
        stalled = in_review and waiting is not None and (waiting > 14 if st == 'Unassigned' else waiting > 30)
        delivery = st in ('25%', '50%', '75%', '100%')
        before = xc is not None and xs is not None and xc < xs
        issue = ('no TA lead' if lead == '' else 'no target date' if xc is None else 'no details/description' if not hd
                 else 'no modality' if r['Modality'] == '' else 'no offer' if r['Primary Programme Offer'] == ''
                 else 'target before start' if before else None)
        phase = (None if st == 'Discontinued' else 'In review' if in_review
                 else 'Overdue' if st != '100%' and past else 'Started & in delivery')
        done = r['Closed'] or r['Resolved']
        month = lambda d: d.month if d is not None and d.year == 2026 else None
        rows.append(dict(
            order=i, office=office or None, region=region, practice=map_practice(r['Global Practice and Cross Sectoral Teams']) or None,
            type=map_type(r['Request Type']) or None, offer=map_offer(r['Primary Programme Offer']) or None, status=st,
            lead=lead or None, reqfor=r['Requested For'], short=r['Short description'], xc=xc, offer_raw=r['Primary Programme Offer'],
            modality=r['Modality'], hd=hd, ph=ph, co=is_co, perf=is_co and st != 'Discontinued', overdue=overdue,
            bucket=None if not overdue else '1–30 days' if late <= 30 else '31–60 days' if late <= 60 else '>60 days',
            active=active, ontrack=active and not past,
            recv=received is not None and today - 30 <= oad(received) <= today,
            opened_m=month(r['Opened']), completed_m=month(done) if st == '100%' else None,
            in_review=in_review, waiting=waiting if in_review else None, stalled=stalled,
            setup=None if not in_review or waiting is None else '0–14 days' if waiting <= 14 else '15–30 days' if waiting <= 30 else '30+ days',
            ready=st == '0%' and hd and lead != '' and xc is not None, nolead0=st == '0%' and lead == '',
            delivery=delivery, passes=delivery and issue is None, before=before,
            due30=not closedish and xc is not None and today <= xc <= today + 30,
            notclosed=closedish and r['Closed'] is None, phase=phase))
    seen = {}
    for r in rows:                          # Possible duplicate
        r['dup'] = False
        if r['co'] and r['reqfor'] != '' and r['short'] != '':
            k = (r['reqfor'] + '|' + r['short']).lower()
            r['dup'] = k in seen
            seen.setdefault(k, r['order'])
    return rows, today


def strict_pct(n, d):
    return None if not d else 100 if n >= d else min(99, round(n / d * 100 + 1e-9))


def performance(rows):
    F = [r for r in rows if r['perf']]
    cnt = lambda f: sum(1 for r in F if f(r))
    leads = {}
    for r in F:
        if r['lead']:
            leads[r['lead']] = leads.get(r['lead'], 0) + 1
    top = max(leads.values())
    return {
        'Total requests': len(F), 'Big ticket': cnt(lambda r: r['type'] == 'Big Ticket'), 'Routine': cnt(lambda r: r['type'] == 'Routine'),
        'Received last 30 days': cnt(lambda r: r['recv']), 'Active & on track': cnt(lambda r: r['ontrack']),
        'Completed': cnt(lambda r: r['status'] == '100%'), 'Overdue': cnt(lambda r: r['overdue']),
        'Overdue 1–30 / 31–60 / >60': '/'.join(str(cnt(lambda r, b=b: r['bucket'] == b)) for b in ('1–30 days', '31–60 days', '>60 days')),
        'Opened by month (Apr→)': '/'.join(str(cnt(lambda r, m=m: r['opened_m'] == m)) for m in range(4, 11)),
        'Completed by month (Apr→)': '/'.join(str(cnt(lambda r, m=m: r['completed_m'] == m)) for m in range(4, 11)),
        'TA leads': len(leads), 'Lead load min / avg / max': f'{min(leads.values())} / {sum(leads.values()) / len(leads):.1f} / {top}',
        'Busiest TA lead': sorted(k for k, v in leads.items() if v == top)[0],
    }


def data_quality(rows):
    F = [r for r in rows if r['co']]
    cnt = lambda f: sum(1 for r in F if f(r))
    d = cnt(lambda r: r['delivery'])
    comp = {
        'TA lead': cnt(lambda r: r['delivery'] and r['lead']),
        'Expected completion': cnt(lambda r: r['delivery'] and r['xc'] is not None),
        'Details/Description': cnt(lambda r: r['delivery'] and r['hd']),
        'Modality': cnt(lambda r: r['delivery'] and r['modality'] != ''),
        'Programme offer': cnt(lambda r: r['delivery'] and r['offer_raw'] != ''),
    }
    return {
        'Awaiting assignment': cnt(lambda r: r['status'] == 'Unassigned'), 'In review (0%)': cnt(lambda r: r['status'] == '0%'),
        'Stalled in setup': cnt(lambda r: r['stalled']), 'In delivery (25%+)': cnt(lambda r: r['status'] in ('25%', '50%', '75%')),
        'Needing cleanup': cnt(lambda r: r['delivery'] and not r['passes']), 'Overdue': cnt(lambda r: r['overdue']),
        'Time in setup 0–14 / 15–30 / 30+': '/'.join(str(cnt(lambda r, b=b: r['setup'] == b)) for b in ('0–14 days', '15–30 days', '30+ days')),
        'Ready to advance': cnt(lambda r: r['ready']), '0% without TA lead': cnt(lambda r: r['nolead0']),
        'Completeness % (lead/date/desc/modality/offer)': '/'.join(str(strict_pct(v, d)) for v in comp.values()),
        'Record quality score %': strict_pct(cnt(lambda r: r['passes']), d),
        'No TA lead / no target / target before start (25%+)': f"{cnt(lambda r: r['delivery'] and not r['lead'])} / "
            f"{cnt(lambda r: r['delivery'] and r['xc'] is None)} / {cnt(lambda r: r['delivery'] and r['before'])}",
        'Placeholder descriptions': cnt(lambda r: r['ph'] and r['status'] != 'Discontinued'),
        'Possible duplicates': cnt(lambda r: r['dup']), 'Due in next 30 days': cnt(lambda r: r['due30']),
        'Completed/discontinued not closed': cnt(lambda r: r['notclosed']),
        'Phase: review / delivery / overdue': '/'.join(str(cnt(lambda r, p=p: r['phase'] == p)) for p in ('In review', 'Started & in delivery', 'Overdue')),
        'Requests with no office (whole export)': sum(1 for r in rows if r['office'] is None),
    }


def survey(path, raw):                      # Survey.m
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def sheet(name):
        ws = wb[name]; ws.reset_dimensions()
        rows = list(ws.iter_rows(values_only=True))
        filled = lambda r: sum(1 for c in r if c not in (None, ''))
        hi = next(i for i, r in enumerate(rows) if filled(r) > 4)
        head = [mtext(h).strip() or f'Column{j + 1}' for j, h in enumerate(rows[hi])]
        return [dict(zip(head, r)) for r in rows[hi + 1:] if filled(r) > 0]
    rawc = {collapse(r.get('ID')): collapse(r.get('Open comment')) for r in sheet('Response data')}
    cases = {}
    for r in raw:
        cases[r['Number']] = r
    scales = {'Satisfaction': ['Very dissatisfied', 'Dissatisfied', 'Neither satisfied nor dissatisfied', 'Satisfied', 'Very satisfied'],
              'Quality': ['Poor', 'Fair', 'Good', 'Very Good', 'Excellent'],
              'Timeliness': ['Much too late', 'Too late', 'Acceptable', 'Timely', 'Very Timely']}
    score = lambda sc, lab: scales[sc].index(lab) + 1 if lab in scales[sc] else None
    nonsub = {'', 'n/a', 'na', 'no comment', 'no comments', 'none', 'nil', '.', '-', '..', '...'}
    out = []
    for r in sheet('Merged data'):
        rid = collapse(r.get('ID'))
        if rid in EXCL_RESP:
            continue
        case_no = collapse(r.get('TA Case Number'))
        body = rawc.get(rid) or collapse(r.get('Open comment'))
        man = MANUAL.get(rid)
        pos, imp, flag = (man if man else [collapse(r.get('Positive Themes')), collapse(r.get('Improvement Theme')), collapse(r.get('Data-quality Flag'))])
        lab = collapse(r.get('Satisfaction'))
        if case_no in SATFIX and lab == SATFIX[case_no][0]:
            lab = SATFIX[case_no][1]
        try:
            rec = float(collapse(r.get('Recommend 0–10')))
        except ValueError:
            rec = None
        case = cases.get(case_no) if case_no else None
        office_raw = (case['Office/Division'] if case and case['Office/Division'] else collapse(r.get('Country Office')))
        office, region = map_office(office_raw)
        survey_office = HQ.get(collapse(r.get('Country Office')), collapse(r.get('Country Office')))
        out.append(dict(id=rid, office=survey_office, region=region, matched=case is not None,
                        type=(map_type(case['Request Type']) or 'Unclassified') if case else 'Unclassified',
                        practice=map_practice(case['Global Practice and Cross Sectoral Teams']) if case else None,
                        written=body != '', substantive=body != '' and body.strip().lower().rstrip('.!?') not in nonsub,
                        pos=[t.strip() for t in pos.split(';') if t.strip()], imp=imp or None, flag=flag or None,
                        sat=score('Satisfaction', lab), qual=score('Quality', collapse(r.get('Quality'))),
                        time=score('Timeliness', collapse(r.get('Timeliness'))), rec=rec,
                        mapped=collapse(r.get('Country Office')) in COORDS))
    return out


def feedback(rs):
    avg = lambda k: round(sum(r[k] for r in rs if r[k] is not None) / sum(1 for r in rs if r[k] is not None), 2)
    pos = {}
    for r in rs:
        for t in r['pos']:
            pos[t] = pos.get(t, 0) + 1
    imp = {}
    for r in rs:
        if r['imp']:
            imp[r['imp']] = imp.get(r['imp'], 0) + 1
    types = {}
    for r in rs:
        types[r['type']] = types.get(r['type'], 0) + 1
    return {
        'Responses': len(rs), 'Written comments': sum(r['written'] for r in rs), 'Substantive comments': sum(r['substantive'] for r in rs),
        'Improvement opportunities': sum(1 for r in rs if r['imp']), 'Data-quality flags': sum(1 for r in rs if r['flag']),
        'Average satisfaction / quality / timeliness': f"{avg('sat'):.2f} / {avg('qual'):.2f} / {avg('time'):.2f}",
        'Average recommend': round(avg('rec'), 1), 'Rated responses': sum(1 for r in rs if r['sat'] is not None),
        'Offices responding': len({r['office'] for r in rs if r['office']}),
        'Matched to a request': sum(r['matched'] for r in rs),
        'Pieces of positive feedback': sum(pos.values()), 'Comments with positive feedback': sum(1 for r in rs if r['pos']),
        'Routine / Big Ticket / Unclassified': f"{types.get('Routine', 0)} / {types.get('Big Ticket', 0)} / {types.get('Unclassified', 0)}",
        'Top positive theme': max(sorted(pos), key=lambda k: pos[k]),
    }


def web_view(w):
    k = dict(w['kpis'])
    dq = dict(w['dq']['kpis'])
    i = lambda s: int(str(s).replace(',', ''))
    big, routine = k['Big ticket vs. routine'].split('/')
    comp = dict(w['dq']['completeness'])
    stack = w['dq']['stack']
    f = w['feedback']
    return {
        'Total requests': i(k['Total requests']), 'Big ticket': i(big), 'Routine': i(routine),
        'Received last 30 days': i(k['Received last 30 days']), 'Active & on track': i(k['Active & on track']),
        'Completed': i(k['Completed']), 'Overdue': i(k['Overdue']),
        'Overdue 1–30 / 31–60 / >60': '/'.join(str(b[1]) for b in w['buckets']),
        'Opened by month (Apr→)': '/'.join(str(m[1]) for m in w['io']),
        'Completed by month (Apr→)': '/'.join(str(m[2]) for m in w['io']),
        'TA leads': w['leads'][0], 'Lead load min / avg / max': f"{w['leads'][1]} / {w['leads'][2]} / {w['leads'][3]}",
        'Busiest TA lead': w['leads'][4],
        'Awaiting assignment': i(dq['Awaiting assignment']), 'In review (0%)': i(dq['In review (0%)']),
        'Stalled in setup': i(dq['Stalled in setup']), 'In delivery (25%+)': i(dq['In delivery (25%+)']),
        'Needing cleanup': i(dq['Needing cleanup']), 'DQ Overdue': i(dq['Overdue']),
        'Time in setup 0–14 / 15–30 / 30+': '/'.join(str(b[1]) for b in w['dq']['aging']),
        'Ready to advance': i(w['dq']['ready'].split(' of ')[0]), '0% without TA lead': i(w['dq']['contradictions'][0]),
        'Completeness % (lead/date/desc/modality/offer)': '/'.join(str(comp[c]) for c in ('TA lead', 'Expected completion', 'Details/Description', 'Modality', 'Programme offer')),
        'Record quality score %': i(w['dq']['score'].rstrip('%')),
        'No TA lead / no target / target before start (25%+)': ' / '.join(str(i(x[1])) for x in w['dq']['flags']),
        'Placeholder descriptions': i(w['dq']['ph']), 'Possible duplicates': i(w['dq']['dup']),
        'Due in next 30 days': i(w['dq']['atRisk']), 'Completed/discontinued not closed': i(w['dq']['notClosed']),
        'Phase: review / delivery / overdue': '/'.join(str(sum(r[j] for r in stack)) for j in (1, 2, 3)),
        'Requests with no office (whole export)': i(w['coUnassigned']),
        'Responses': f['kpi']['responses'], 'Written comments': f['kpi']['written'], 'Substantive comments': f['kpi']['substantive'],
        'Improvement opportunities': f['kpi']['improvement'], 'Data-quality flags': f['kpi']['flags'],
        'Average satisfaction / quality / timeliness': f"{f['avg']['sat']:.2f} / {f['avg']['qual']:.2f} / {f['avg']['time']:.2f}",
        'Average recommend': round(f['avg']['rec'], 1), 'Rated responses': f['rated'],
        'Routine / Big Ticket / Unclassified': f"{f['types'].get('Routine', 0)} / {f['types'].get('Big Ticket', 0)} / {f['types'].get('Unclassified', 0)}",
        'Pieces of positive feedback': sum(x['n'] for x in f['positive']), 'Top positive theme': f['positive'][0]['label'],
    }


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__.strip())
    raw = load_requests(sys.argv[1])
    rows, today = requests_table(raw)
    perf, dq, fb = performance(rows), data_quality(rows), feedback(survey(sys.argv[2], raw))
    dq['DQ Overdue'] = dq.pop('Overdue')
    web = web_view(json.load(open(sys.argv[3]))) if len(sys.argv) > 3 else {}
    print(f'As of {(EPOCH + datetime.timedelta(days=today)).strftime("%d %b %Y")}\n')
    mism = 0
    for title, block in (('PERFORMANCE', perf), ('DATA QUALITY', dq), ('FEEDBACK', fb)):
        print(title)
        for k, v in block.items():
            w = web.get(k, '')
            flag = '' if not web or k not in web else ('  ok' if str(w) == str(v) else '  << DIFFERS')
            mism += flag.endswith('DIFFERS')
            print(f'  {k:<52} {str(v):<24} {str(w):<24}{flag}')
        print()
    if web:
        print(f'{mism} difference(s) against the web dashboard')


if __name__ == '__main__':
    main()
